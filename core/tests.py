from io import BytesIO
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from PIL import Image

from .models import ProcessingJob
from .services import ArchiveConverter, ConversionError, ConversionRouter


def image_upload(name="sample.png", mode="RGB", color=(200, 40, 40, 255)):
    buffer = BytesIO()
    image = Image.new(mode, (16, 16), color)
    image.save(buffer, format="PNG")
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")


@override_settings(MEDIA_ROOT="D:/tmp/thuna-test-media")
class ConversionServiceTests(TestCase):
    def make_job(self, upload, operation="convert"):
        return ProcessingJob.objects.create(
            session_key="tests",
            operation=operation,
            original=upload,
            original_name=upload.name,
            original_size=upload.size,
        )

    def test_png_converts_to_webp_and_validates_output(self):
        job = self.make_job(image_upload())

        ConversionRouter().convert(job, "webp", 80)
        job.refresh_from_db()

        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "webp")
        self.assertGreater(job.result_size, 0)
        with Image.open(job.result.path) as converted:
            self.assertEqual(converted.format, "WEBP")

    def test_transparent_png_converts_to_jpeg_without_crashing(self):
        job = self.make_job(image_upload(mode="RGBA", color=(0, 120, 255, 80)))

        ConversionRouter().convert(job, "jpg", 80)
        job.refresh_from_db()

        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        with Image.open(job.result.path) as converted:
            self.assertEqual(converted.format, "JPEG")
            self.assertEqual(converted.mode, "RGB")

    def test_invalid_conversion_is_rejected(self):
        job = self.make_job(image_upload())

        with self.assertRaises(ConversionError) as raised:
            ConversionRouter().convert(job, "mp3", 80)

        self.assertEqual(raised.exception.code, "UNSUPPORTED_FORMAT")

    def test_zip_create_handles_duplicate_names(self):
        uploads = [
            SimpleUploadedFile("photo.jpg", b"first"),
            SimpleUploadedFile("photo.jpg", b"second"),
        ]

        result = ArchiveConverter("D:/tmp").create_zip(uploads, "download")

        with ZipFile(BytesIO(result.data)) as archive:
            self.assertEqual(archive.namelist(), ["photo.jpg", "photo-2.jpg"])

    def test_zip_slip_entry_is_rejected(self):
        buffer = BytesIO()
        with ZipFile(buffer, "w") as archive:
            archive.writestr("../evil.txt", b"bad")

        with self.assertRaises(ConversionError) as raised:
            ArchiveConverter("D:/tmp").inspect_zip(buffer.getvalue())

        self.assertEqual(raised.exception.code, "UNSAFE_ARCHIVE")

    def test_docx_converts_to_pdf_natively(self):
        import docx
        doc = docx.Document()
        doc.add_heading("Test Heading", 1)
        doc.add_paragraph("This is a test paragraph inside docx.")
        buf = BytesIO()
        doc.save(buf)
        upload = SimpleUploadedFile("sample.docx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "pdf", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "pdf")
        self.assertGreater(job.result_size, 0)

    def test_docx_converts_to_txt_natively(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        buf = BytesIO()
        doc.save(buf)
        upload = SimpleUploadedFile("sample.docx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "txt", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertIn(b"First paragraph.", job.result.read())

    def test_xlsx_converts_to_csv_natively(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Age", "Role"])
        ws.append(["Alice", 30, "Engineer"])
        buf = BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile("data.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "csv", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        content = job.result.read().decode("utf-8")
        self.assertIn("Alice", content)
        self.assertIn("Engineer", content)

    def test_xlsx_converts_to_pdf_natively(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Col1", "Col2"])
        ws.append(["Val1", "Val2"])
        buf = BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile("data.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "pdf", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "pdf")

    def test_pptx_converts_to_pdf_natively(self):
        import pptx
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Presentation Test"
        buf = BytesIO()
        prs.save(buf)
        upload = SimpleUploadedFile("deck.pptx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "pdf", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "pdf")

    def test_process_merge_mixed_formats(self):
        import docx
        from .services import process_merge
        # 1. Image
        img_upload = image_upload("img1.png")
        # 2. DOCX
        doc = docx.Document()
        doc.add_paragraph("DOCX page to merge.")
        buf = BytesIO()
        doc.save(buf)
        docx_upload = SimpleUploadedFile("doc1.docx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        job = self.make_job(img_upload, operation="merge")
        result = process_merge(job, [img_upload, docx_upload], "pdf", 80)
        self.assertEqual(result.output_format, "pdf")
        self.assertGreater(len(result.data), 0)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)

    def test_pdf_to_docx_preserves_content(self):
        import fitz
        import docx
        # Create a clean test PDF with multiple structured paragraphs
        doc_pdf = fitz.open()
        page = doc_pdf.new_page(width=595, height=842)
        page.insert_text((50, 80), "Header Title for Document", fontname="helv", fontsize=18)
        page.insert_text((50, 130), "Paragraph 1: The quick brown fox jumps over the lazy dog.", fontname="helv", fontsize=11)
        page.insert_text((50, 160), "Paragraph 2: Numerical data 12345.67 and special characters test.", fontname="helv", fontsize=11)
        buf = BytesIO()
        doc_pdf.save(buf)
        doc_pdf.close()

        upload = SimpleUploadedFile("sample.pdf", buf.getvalue(), content_type="application/pdf")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "docx", 80)
        job.refresh_from_db()

        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "docx")

        # Verify content in the output DOCX
        doc_result = docx.Document(job.result.path)
        all_text = "\n".join(p.text for p in doc_result.paragraphs)
        self.assertIn("Header Title for Document", all_text)
        self.assertIn("quick brown fox jumps over the lazy dog", all_text)
        self.assertIn("Numerical data 12345.67", all_text)

    def test_pdf_to_doc_format(self):
        import fitz
        doc_pdf = fitz.open()
        page = doc_pdf.new_page(width=595, height=842)
        page.insert_text((50, 80), "DOC format test.", fontname="helv", fontsize=14)
        buf = BytesIO()
        doc_pdf.save(buf)
        doc_pdf.close()

        upload = SimpleUploadedFile("sample.pdf", buf.getvalue(), content_type="application/pdf")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "doc", 80)
        job.refresh_from_db()

        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "doc")
        self.assertTrue(job.result.name.endswith(".doc"))

    def test_500mb_upload_validation(self):
        from .forms import validate_upload
        from django import forms
        # Under 500 MB mock file (e.g. 400 MB reported size)
        valid_mock = SimpleUploadedFile("valid.pdf", b"%PDF-1.4 test data", content_type="application/pdf")
        valid_mock.size = 400 * 1024 * 1024
        # Validate should not raise a size error
        validate_upload(valid_mock)

        # Over 500 MB mock file (501 MB)
        oversized = SimpleUploadedFile("big.pdf", b"%PDF-1.4 test data", content_type="application/pdf")
        oversized.size = 501 * 1024 * 1024
        with self.assertRaises(forms.ValidationError) as ctx:
            validate_upload(oversized)
        self.assertIn("500 MB", str(ctx.exception))

    def test_privacy_page_renders_cleanly(self):
        from django.test import Client
        client = Client(HTTP_HOST="127.0.0.1")
        response = client.get("/privacy/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Privacy Policy")
        self.assertContains(response, "Our Core Privacy Promise")
        self.assertContains(response, "cta-banner")

    def test_terms_page_renders_cleanly(self):
        from django.test import Client
        client = Client(HTTP_HOST="127.0.0.1")
        response = client.get("/terms/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Terms of Service")
        self.assertContains(response, "500 MB")
        self.assertContains(response, "cta-banner")

    def test_robots_txt_served(self):
        from django.test import Client
        client = Client(HTTP_HOST="127.0.0.1")
        response = client.get("/robots.txt")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/plain")
        self.assertIn("Disallow: /jobs/", response.content.decode("utf-8"))
        self.assertIn("Allow: /", response.content.decode("utf-8"))

    def test_custom_404_view(self):
        from django.test import Client
        client = Client(HTTP_HOST="127.0.0.1")
        response = client.get("/non-existent-page-url-12345/")
        self.assertEqual(response.status_code, 404)
        self.assertContains(response, "404", status_code=404)
        self.assertContains(response, "Return to Home", status_code=404)

    def test_seo_meta_tags_and_cookie_banner(self):
        from django.test import Client
        client = Client(HTTP_HOST="127.0.0.1")
        response = client.get("/")
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        # SEO titles & meta description
        self.assertIn("<title>THUNA - Free Online File Converter, Compressor & PDF Tools</title>", content)
        self.assertIn('<meta name="description"', content)
        self.assertIn('<meta property="og:title"', content)
        # Cookie banner & alt text
        self.assertIn('id="cookieConsentBanner"', content)
        self.assertIn('id="btnAcceptCookies"', content)
        self.assertIn('alt="THUNA brand logo - Free File Tools"', content)

    def test_capabilities_detect_ffmpeg_and_ffprobe(self):
        from .services import get_capabilities
        caps = get_capabilities()
        self.assertTrue(caps["video"])
        self.assertTrue(caps["audio"])
        self.assertTrue(caps["pdf"])
        self.assertTrue(caps["image"])

    def test_audio_conversion_and_compression(self):
        import shutil
        import subprocess
        ffmpeg = shutil.which("ffmpeg")
        self.assertIsNotNone(ffmpeg)

        wav_buf = subprocess.run(
            [ffmpeg, "-y", "-f", "lavfi", "-i", "sine=frequency=1000:duration=1", "-c:a", "pcm_s16le", "-f", "wav", "pipe:1"],
            capture_output=True, check=True
        ).stdout

        upload = SimpleUploadedFile("sample.wav", wav_buf, content_type="audio/wav")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "mp3", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "mp3")
        self.assertGreater(job.result_size, 0)

        # Audio compression
        cmp_job = self.make_job(SimpleUploadedFile("audio.mp3", job.result.read(), content_type="audio/mpeg"), operation="compress")
        ConversionRouter().convert(cmp_job, "mp3", 50)
        cmp_job.refresh_from_db()
        self.assertEqual(cmp_job.status, ProcessingJob.Status.SUCCESS)
        self.assertGreater(cmp_job.result_size, 0)

    def test_video_conversion_and_compression(self):
        import shutil
        import subprocess
        ffmpeg = shutil.which("ffmpeg")
        self.assertIsNotNone(ffmpeg)

        mp4_buf = subprocess.run(
            [ffmpeg, "-y", "-f", "lavfi", "-i", "testsrc=duration=1:size=160x120:rate=10",
             "-f", "lavfi", "-i", "sine=frequency=1000:duration=1",
             "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-f", "mp4", "-movflags", "+frag_keyframe+empty_moov", "pipe:1"],
            capture_output=True, check=True
        ).stdout

        upload = SimpleUploadedFile("clip.mp4", mp4_buf, content_type="video/mp4")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "webm", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(job.output_format, "webm")

        # Extract audio from video
        ext_job = self.make_job(SimpleUploadedFile("clip.mp4", mp4_buf, content_type="video/mp4"))
        ConversionRouter().convert(ext_job, "mp3", 80)
        ext_job.refresh_from_db()
        self.assertEqual(ext_job.status, ProcessingJob.Status.SUCCESS)
        self.assertEqual(ext_job.output_format, "mp3")

        # Video compression
        cmp_job = self.make_job(SimpleUploadedFile("clip.mp4", mp4_buf, content_type="video/mp4"), operation="compress")
        ConversionRouter().convert(cmp_job, "mp4", 50)
        cmp_job.refresh_from_db()
        self.assertEqual(cmp_job.status, ProcessingJob.Status.SUCCESS)
        self.assertGreater(cmp_job.result_size, 0)

    def test_rtf_conversion_fallback(self):
        rtf_data = rb"{\rtf1\ansi\deff0 {\fonttbl {\f0 Times New Roman;}}\f0\fs24 Hello \b RTF\b0!}"
        upload = SimpleUploadedFile("sample.rtf", rtf_data, content_type="application/rtf")
        job = self.make_job(upload)
        ConversionRouter().convert(job, "txt", 80)
        job.refresh_from_db()
        self.assertEqual(job.status, ProcessingJob.Status.SUCCESS)
        self.assertIn(b"Hello RTF!", job.result.read())

    def test_gitignore_covers_secrets_and_environment(self):
        from pathlib import Path
        gitignore_path = Path(__file__).resolve().parent.parent / ".gitignore"
        self.assertTrue(gitignore_path.exists())
        content = gitignore_path.read_text(encoding="utf-8")
        self.assertIn(".env", content)
        self.assertIn("*.sqlite3", content)
        self.assertIn("media/*", content)
        self.assertIn(".venv/", content)




