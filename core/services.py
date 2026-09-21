import csv
import html
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone
from PIL import Image, ImageOps, UnidentifiedImageError

from .models import ProcessingJob

Image.MAX_IMAGE_PIXELS = getattr(settings, "MAX_IMAGE_PIXELS", 80_000_000)

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
    HEIF_AVAILABLE = True
except ImportError:
    HEIF_AVAILABLE = False

try:
    import static_ffmpeg

    static_ffmpeg.add_paths()
except Exception:
    pass



class ConversionError(Exception):
    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


@dataclass
class ConversionResult:
    name: str
    data: bytes
    output_format: str
    metadata: dict | None = None


IMAGE_OUTPUT_FORMATS = {
    "jpg": "JPEG",
    "jpeg": "JPEG",
    "png": "PNG",
    "webp": "WEBP",
    "bmp": "BMP",
    "tif": "TIFF",
    "tiff": "TIFF",
    "gif": "GIF",
    "pdf": "PDF",
}
IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "bmp", "gif", "tif", "tiff", "heic", "heif"}
PDF_EXTENSIONS = {"pdf"}
OFFICE_EXTENSIONS = {"doc", "docx", "odt", "rtf", "ppt", "pptx", "xls", "xlsx"}
TEXT_EXTENSIONS = {"txt", "csv", "html", "htm", "md", "markdown"}
VIDEO_EXTENSIONS = {"mp4", "mov", "avi", "mkv", "webm"}
AUDIO_EXTENSIONS = {"mp3", "wav", "m4a", "aac", "flac", "ogg"}
ARCHIVE_EXTENSIONS = {"zip"}

COMPATIBILITY = {
    "pdf": {"docx", "doc", "txt", "jpg", "jpeg", "png", "html"},
    "image": {"jpg", "jpeg", "png", "webp", "bmp", "tiff", "pdf"},
    "office": {"pdf", "txt", "docx", "doc", "csv", "xlsx", "html"},
    "text": {"pdf", "docx", "html", "txt", "xlsx", "csv"},
    "video": {"mp4", "mov", "avi", "mkv", "webm", "mp3", "wav"},
    "audio": {"mp3", "wav", "m4a", "aac", "flac", "ogg"},
    "archive": {"extract", "zip"},
}

ERROR_MESSAGES = {
    "UNSUPPORTED_FORMAT": "This conversion is not supported.",
    "INVALID_FILE": "The uploaded file type does not match its contents.",
    "CORRUPTED_FILE": "The uploaded file appears corrupted.",
    "MISSING_ENGINE": "The required conversion engine is not installed on this server.",
    "CONVERSION_FAILED": "The file could not be converted.",
    "OUTPUT_VALIDATION_FAILED": "The converted file failed validation.",
    "NO_AUDIO_STREAM": "This video does not contain an audio track.",
    "NO_VIDEO_STREAM": "This file does not contain a video stream.",
    "OCR_REQUIRED": "This PDF has no selectable text and OCR is unavailable.",
    "ARCHIVE_TOO_LARGE": "The archive is too large to process safely.",
    "UNSAFE_ARCHIVE": "The ZIP archive contains unsafe file paths.",
    "TIMEOUT": "The conversion timed out.",
}


def setting(name, default):
    return getattr(settings, name, default)


def binary_setting(name, default):
    return os.environ.get(name, getattr(settings, name, default))


def normalize_ext(value):
    return (value or "").lower().strip().lstrip(".")


def safe_stem(filename):
    stem = Path(filename).stem or "file"
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip(".-")
    return stem[:80] or "file"


def read_upload(uploaded):
    uploaded.open("rb")
    uploaded.seek(0)
    data = uploaded.read()
    uploaded.seek(0)
    return data


def has_binary(binary):
    if shutil.which(binary) is not None or Path(binary).exists():
        return True
    if binary in ("ffmpeg", "ffprobe"):
        try:
            import static_ffmpeg

            static_ffmpeg.add_paths()
            return shutil.which(binary) is not None
        except Exception:
            pass
    return False



def get_capabilities():
    return {
        "image": True,
        "pdf": module_available("fitz"),
        "office": (module_available("docx") and module_available("openpyxl"))
        or has_binary(binary_setting("LIBREOFFICE_BINARY", "libreoffice"))
        or has_binary(binary_setting("LIBREOFFICE_BINARY", "soffice")),
        "video": has_binary(binary_setting("FFMPEG_BINARY", "ffmpeg"))
        and has_binary(binary_setting("FFPROBE_BINARY", "ffprobe")),
        "audio": has_binary(binary_setting("FFMPEG_BINARY", "ffmpeg"))
        and has_binary(binary_setting("FFPROBE_BINARY", "ffprobe")),
        "zip": True,
        "ocr": has_binary(binary_setting("TESSERACT_BINARY", "tesseract")) and module_available("pytesseract"),
        "heic": HEIF_AVAILABLE,
        "celery": module_available("celery"),
        "redis": module_available("redis"),
    }


def module_available(name):
    if name == "fitz":
        name = "pymupdf"
    try:
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            __import__(name)
        return True
    except ImportError:
        if name == "pymupdf":
            try:
                import fitz
                return True
            except ImportError:
                return False
        return False


def get_fitz():
    try:
        import pymupdf as fitz
        return fitz
    except ImportError:
        try:
            import fitz
            return fitz
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "PDF processing requires PyMuPDF.") from exc


def detect_file(uploaded):
    ext = normalize_ext(Path(uploaded.name).suffix)
    header = read_upload(uploaded)[:16]
    if ext in IMAGE_EXTENSIONS:
        if ext in {"heic", "heif"} and not HEIF_AVAILABLE:
            raise ConversionError("MISSING_ENGINE", "HEIC conversion requires HEIF support.")
        try:
            uploaded.open("rb")
            image = Image.open(uploaded)
            image.verify()
            fmt = (image.format or ext).lower()
            uploaded.seek(0)
            return "image", "jpg" if fmt == "jpeg" else fmt
        except Exception as exc:
            raise ConversionError("CORRUPTED_FILE", "This image file appears unreadable or corrupted.") from exc
    if ext == "pdf":
        if not header.startswith(b"%PDF-"):
            raise ConversionError("INVALID_FILE", "This does not appear to be a valid PDF.")
        return "pdf", "pdf"
    if ext == "zip":
        try:
            uploaded.open("rb")
            with ZipFile(uploaded) as archive:
                if archive.testzip() is not None:
                    raise ConversionError("CORRUPTED_FILE", "This ZIP file appears corrupted.")
            uploaded.seek(0)
            return "archive", "zip"
        except BadZipFile as exc:
            raise ConversionError("INVALID_FILE", "This does not appear to be a valid ZIP file.") from exc
    if ext in OFFICE_EXTENSIONS:
        if ext in {"docx", "pptx", "xlsx", "odt"}:
            if not header.startswith(b"PK"):
                raise ConversionError("INVALID_FILE", f"This does not appear to be a valid {ext.upper()} file.")
            return "office", ext
        if ext in {"doc", "ppt", "xls"}:
            if header.startswith(b"PK"):
                mapped = {"doc": "docx", "ppt": "pptx", "xls": "xlsx"}.get(ext, ext)
                return "office", mapped
            if not header.startswith(b"\xd0\xcf\x11\xe0"):
                raise ConversionError("INVALID_FILE", f"This does not appear to be a valid {ext.upper()} file.")
            return "office", ext
        if ext == "rtf":
            if not header.startswith(b"{\\rtf"):
                raise ConversionError("INVALID_FILE", "This does not appear to be a valid RTF file.")
            return "office", "rtf"
        return "office", ext
    if ext in TEXT_EXTENSIONS:
        return "text", ext
    if ext in VIDEO_EXTENSIONS:
        return "video", ext
    if ext in AUDIO_EXTENSIONS:
        return "audio", ext
    raise ConversionError("UNSUPPORTED_FORMAT", f"THUNA does not support {ext.upper() or 'this'} files.")


def validate_requested(category, source_format, target_format, operation):
    target = normalize_ext(target_format)
    if operation == "compress":
        return target or source_format
    if operation == "zip_create":
        return "zip"
    if operation == "zip_extract":
        return "extract"
    if not target:
        raise ConversionError("UNSUPPORTED_FORMAT", "Please select an output format for conversion.")
    allowed = COMPATIBILITY.get(category, set())
    if target not in allowed:
        raise ConversionError(
            "UNSUPPORTED_FORMAT",
            f"{source_format.upper()} to {target.upper()} is not a valid conversion.",
        )
    return target


class ConversionRouter:
    def convert(self, job, output_format="", quality=80, files=None):
        job.status = ProcessingJob.Status.PROCESSING
        job.save(update_fields=["status"])

        with tempfile.TemporaryDirectory(prefix=f"thuna-{job.token}-", dir=setting("CONVERSION_TMP_ROOT", None)) as workdir:
            job.temporary_path = workdir
            job.save(update_fields=["temporary_path"])
            category, source_format = detect_file(job.original)
            target = validate_requested(category, source_format, output_format, job.operation)
            job.input_format = source_format
            job.output_format = target
            job.save(update_fields=["input_format", "output_format"])

            if job.operation == "compress" and category == "archive":
                raise ConversionError("UNSUPPORTED_FORMAT", "ZIP compression has been removed. Choose an image or PDF to compress.")
            elif job.operation == "zip_extract":
                result = ArchiveConverter(workdir).extract_zip(job)
            elif job.operation == "zip_create":
                result = ArchiveConverter(workdir).create_zip(files or [], safe_stem(job.original_name))
            elif job.operation == "compress" and category == "pdf":
                result = PDFConverter(workdir).compress(job, quality)
            elif job.operation == "compress" and category == "image":
                result = ImageConverter(workdir).convert(job, target, quality)
            elif job.operation == "compress" and category == "office" and source_format == "docx":
                result = DocumentConverter(workdir).compress_docx(job, quality)
            elif job.operation == "compress" and category in {"video", "audio"}:
                result = MediaConverter(workdir).compress(job, source_format, quality)
            elif category == "image":
                result = ImageConverter(workdir).convert(job, target, quality)
            elif category == "pdf":
                result = PDFConverter(workdir).convert(job, target, quality)
            elif category in {"office", "text"}:
                result = DocumentConverter(workdir).convert(job, source_format, target)
            elif category in {"video", "audio"}:
                result = MediaConverter(workdir).convert(job, source_format, target)
            elif category == "archive":
                result = ArchiveConverter(workdir).extract_zip(job) if target == "extract" else ArchiveConverter(workdir).recompress_zip(job)
            else:
                raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

            validate_output(result.data, result.output_format)
            save_success(job, result)
            return result


class ImageConverter:
    def __init__(self, workdir):
        self.workdir = Path(workdir)

    def convert(self, job, target, quality=80):
        target = "jpg" if target == "jpeg" else target
        if target not in IMAGE_OUTPUT_FORMATS:
            raise ConversionError("UNSUPPORTED_FORMAT", f"Images cannot be exported as {target.upper()}.")
        if target in {"heic", "heif"}:
            raise ConversionError("UNSUPPORTED_FORMAT", "HEIC output is not supported.")

        job.original.open("rb")
        image = Image.open(job.original)
        image = ImageOps.exif_transpose(image)
        if getattr(image, "is_animated", False) and target not in {"gif", "webp"}:
            raise ConversionError("UNSUPPORTED_FORMAT", "Animated conversion is not supported for this format.")

        fmt = IMAGE_OUTPUT_FORMATS[target]
        if getattr(image, "is_animated", False):
            frames = []
            durations = []
            for index in range(getattr(image, "n_frames", 1)):
                image.seek(index)
                frame = ImageOps.exif_transpose(image.copy())
                if target == "webp" and frame.mode not in {"RGB", "RGBA"}:
                    frame = frame.convert("RGBA")
                frames.append(frame)
                durations.append(image.info.get("duration", 100))
            buffer = BytesIO()
            save_args = {
                "format": fmt,
                "save_all": True,
                "append_images": frames[1:],
                "duration": durations,
                "loop": image.info.get("loop", 0),
            }
            if fmt == "WEBP":
                save_args["quality"] = int(quality)
            frames[0].save(buffer, **save_args)
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", buffer.getvalue(), target)

        if fmt in {"JPEG", "PDF"} and image.mode not in ("RGB", "L"):
            background = Image.new("RGB", image.size, "white")
            if image.mode in ("RGBA", "LA"):
                background.paste(image.convert("RGBA"), mask=image.convert("RGBA").split()[-1])
                image = background
            else:
                image = image.convert("RGB")
        elif image.mode == "P":
            image = image.convert("RGBA" if "transparency" in image.info else "RGB")

        buffer = BytesIO()
        save_args = {"format": fmt}
        if fmt in {"JPEG", "WEBP"}:
            save_args.update({"quality": int(quality), "optimize": True})
        elif fmt in {"PNG", "TIFF"}:
            save_args["optimize"] = True
        image.save(buffer, **save_args)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", buffer.getvalue(), target)


class PDFConverter:
    def __init__(self, workdir):
        self.workdir = Path(workdir)

    def _open(self, job):
        fitz = get_fitz()
        return fitz.open(stream=read_upload(job.original), filetype="pdf")

    def convert(self, job, target, quality=80):
        if target in {"jpg", "jpeg", "png"}:
            return self.to_images(job, target, quality)
        if target == "txt":
            return self.to_txt(job)
        if target in {"docx", "doc"}:
            res = self.to_docx(job)
            if target == "doc":
                return ConversionResult(f"{safe_stem(job.original_name)}-thuna.doc", res.data, "doc")
            return res
        if target == "html":
            doc = self._open(job)
            try:
                pages_html = [page.get_text("html") for page in doc]
            finally:
                doc.close()
            full_html = (
                "<!doctype html><html><head><meta charset='utf-8'></head><body>"
                + "".join(pages_html)
                + "</body></html>"
            )
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.html", full_html.encode("utf-8"), "html")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def compress(self, job, quality=80):
        doc = self._open(job)
        buffer = BytesIO()
        try:
            if hasattr(doc, "rewrite_images"):
                doc.rewrite_images(dpi_threshold=200, dpi_target=150, quality=int(quality), lossy=True, lossless=True)
            doc.save(buffer, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
        finally:
            doc.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buffer.getvalue(), "pdf")

    def to_txt(self, job):
        doc = self._open(job)
        try:
            text = "\n\n".join(page.get_text("text") for page in doc)
        finally:
            doc.close()
        if not text.strip():
            text = self.ocr_pdf(job)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", text.encode("utf-8"), "txt")

    def to_docx(self, job):
        # 1. High-fidelity conversion using pdf2docx
        try:
            from pdf2docx import Converter
            input_pdf = self.workdir / f"input_{job.token}.pdf"
            output_docx = self.workdir / f"output_{job.token}.docx"
            input_pdf.write_bytes(read_upload(job.original))
            cv = Converter(str(input_pdf))
            cv.convert(str(output_docx))
            cv.close()
            if output_docx.exists() and output_docx.stat().st_size > 0:
                data = output_docx.read_bytes()
                return ConversionResult(f"{safe_stem(job.original_name)}-thuna.docx", data, "docx")
        except Exception:
            pass

        # 2. Robust fallback preserving sorted layout and unwrapped paragraphs
        try:
            from docx import Document
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "PDF to DOCX requires python-docx.") from exc
        document = Document()
        doc = self._open(job)
        try:
            for page in doc:
                blocks = page.get_text("blocks", sort=True)
                for b in blocks:
                    if len(b) > 4:
                        txt = b[4].strip()
                        if txt:
                            clean_lines = [line.strip() for line in txt.splitlines() if line.strip()]
                            paragraph_text = " ".join(clean_lines)
                            document.add_paragraph(paragraph_text)
        finally:
            doc.close()
        buffer = BytesIO()
        document.save(buffer)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.docx", buffer.getvalue(), "docx")

    def to_images(self, job, target, quality):
        fitz = get_fitz()

        target = "jpg" if target == "jpeg" else target
        doc = self._open(job)
        pages = []
        try:
            for index, page in enumerate(doc):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                buffer = BytesIO()
                save_args = {"format": IMAGE_OUTPUT_FORMATS[target], "optimize": True}
                if target in {"jpg", "webp"}:
                    save_args["quality"] = int(quality)
                image.save(buffer, **save_args)
                pages.append((f"page-{index + 1}.{target}", buffer.getvalue()))
        finally:
            doc.close()

        if len(pages) == 1:
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", pages[0][1], target)
        zip_buffer = BytesIO()
        with ZipFile(zip_buffer, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
            for name, data in pages:
                archive.writestr(name, data)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna-pages.zip", zip_buffer.getvalue(), "zip")

    def ocr_pdf(self, job):
        if not get_capabilities()["ocr"]:
            raise ConversionError("OCR_REQUIRED", ERROR_MESSAGES["OCR_REQUIRED"])
        fitz = get_fitz()
        import pytesseract

        doc = self._open(job)
        parts = []
        try:
            for page in doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                parts.append(pytesseract.image_to_string(image))
        finally:
            doc.close()
        text = "\n\n".join(parts).strip()
        if not text:
            raise ConversionError("OCR_REQUIRED", ERROR_MESSAGES["OCR_REQUIRED"])
        return text


class DocumentConverter:
    def __init__(self, workdir):
        self.workdir = Path(workdir)

    def convert(self, job, source_format, target):
        if source_format == "txt":
            return self.txt(job, target)
        if source_format in {"md", "markdown"}:
            return self.markdown(job, target)
        if source_format in {"html", "htm"}:
            return self.html(job, target)
        if source_format == "csv":
            if target == "xlsx":
                return self.csv_to_xlsx(job)
            if target == "pdf":
                return self.csv_to_pdf(job)
            if target in {"txt", "csv"}:
                return self.txt(job, "txt")
        if source_format == "docx":
            if target == "pdf":
                return self.docx_to_pdf(job)
            if target == "txt":
                return self.docx_to_txt(job)
            if target == "html":
                return self.docx_to_html(job)
            if target in {"docx", "doc"}:
                return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", read_upload(job.original), target)
        if source_format == "doc":
            return self.doc(job, target)
        if source_format == "rtf":
            return self.rtf(job, target)
        if source_format == "xls":
            return self.xls(job, target)
        if source_format in {"xlsx"}:
            if target == "csv":
                return self.xlsx_to_csv(job)
            if target == "pdf":
                return self.xlsx_to_pdf(job)
            if target == "txt":
                return self.xlsx_to_txt(job)
            if target == "html":
                return self.xlsx_to_html(job)
            if target in {"xlsx", "xls"}:
                return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", read_upload(job.original), target)
        if source_format == "ppt":
            return self.ppt(job, target)
        if source_format in {"pptx"}:
            if target == "pdf":
                return self.pptx_to_pdf(job)
            if target == "txt":
                return self.pptx_to_txt(job)
            if target in {"pptx", "ppt"}:
                return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", read_upload(job.original), target)
        if source_format in OFFICE_EXTENSIONS:
            binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
            if has_binary(binary) or has_binary("soffice"):
                return self.libreoffice(job, target)
            raise ConversionError(
                "UNSUPPORTED_FORMAT",
                f"Conversion from {source_format.upper()} to {target.upper()} requires LibreOffice on the server, or converting to PDF/TXT/CSV.",
            )
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def _extract_text_from_ole(self, raw):
        parts = []
        try:
            u16 = raw.decode("utf-16le", errors="ignore")
            u16_clean = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", u16)
            words = [w.strip() for w in u16_clean.split() if len(w.strip()) > 1]
            if len(words) >= 3:
                parts.append(" ".join(words))
        except Exception:
            pass
        ascii_matches = re.findall(rb"[\x20-\x7E\r\n\t]{4,}", raw)
        ascii_text = "\n".join(m.decode("latin-1", errors="ignore").strip() for m in ascii_matches if m.strip())
        if ascii_text:
            parts.append(ascii_text)
        return "\n\n".join(parts).strip()

    def doc(self, job, target):
        binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
        if has_binary(binary) or has_binary("soffice"):
            try:
                return self.libreoffice(job, target)
            except Exception:
                pass
        raw = read_upload(job.original)
        text = ""
        try:
            import olefile

            ole = olefile.OleFileIO(BytesIO(raw))
            if ole.exists("WordDocument"):
                stream = ole.openstream("WordDocument").read()
                text = self._extract_text_from_ole(stream)
        except Exception:
            pass
        if not text:
            text = self._extract_text_from_ole(raw)
        if not text:
            text = f"Document: {job.original_name}"

        if target == "txt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", text.encode("utf-8"), "txt")
        if target in {"docx", "doc"}:
            from docx import Document

            doc = Document()
            for line in text.splitlines():
                if line.strip():
                    doc.add_paragraph(line.strip())
            buf = BytesIO()
            doc.save(buf)
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", buf.getvalue(), target)
        if target == "pdf":
            return text_to_pdf(text, f"{safe_stem(job.original_name)}-thuna.pdf")
        if target == "html":
            body = "<br>\n".join(html.escape(line) for line in text.splitlines() if line.strip())
            return ConversionResult(
                f"{safe_stem(job.original_name)}-thuna.html",
                f"<!doctype html><html><body>{body}</body></html>".encode("utf-8"),
                "html",
            )
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def xls(self, job, target):
        binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
        if has_binary(binary) or has_binary("soffice"):
            try:
                return self.libreoffice(job, target)
            except Exception:
                pass
        try:
            import xlrd

            wb = xlrd.open_workbook(file_contents=read_upload(job.original))
            sheet = wb.sheet_by_index(0)
            rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        except Exception as exc:
            raise ConversionError("CORRUPTED_FILE", "Could not read XLS workbook.") from exc

        if target == "csv":
            out = StringIO()
            writer = csv.writer(out, lineterminator="\n")
            for row in rows:
                if any(c is not None and str(c).strip() != "" for c in row):
                    writer.writerow([str(c) if c is not None else "" for c in row])
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.csv", out.getvalue().encode("utf-8"), "csv")
        if target in {"xlsx", "xls"}:
            from openpyxl import Workbook

            wb_out = Workbook()
            ws = wb_out.active
            for row in rows:
                ws.append(row)
            buf = BytesIO()
            wb_out.save(buf)
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.xlsx", buf.getvalue(), "xlsx")
        if target == "txt":
            lines = ["\t".join(str(c) for c in row if c is not None) for row in rows]
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", "\n".join(lines).encode("utf-8"), "txt")
        if target == "html":
            body = ["<table border='1' cellpadding='6' cellspacing='0'>"]
            for idx, row in enumerate(rows):
                tag = "th" if idx == 0 else "td"
                cells = "".join(f"<{tag}>{html.escape(str(c) if c is not None else '')}</{tag}>" for c in row)
                body.append(f"<tr>{cells}</tr>")
            body.append("</table>")
            return ConversionResult(
                f"{safe_stem(job.original_name)}-thuna.html",
                f"<!doctype html><html><body>{''.join(body)}</body></html>".encode("utf-8"),
                "html",
            )
        if target == "pdf":
            fitz = get_fitz()

            if not rows:
                raise ConversionError("CONVERSION_FAILED", "Worksheet contains no readable rows.")
            num_cols = max(len(r) for r in rows)
            is_landscape = num_cols > 6
            page_w = 842.0 if is_landscape else 595.0
            page_h = 595.0 if is_landscape else 842.0
            margin_x, margin_top, margin_bottom = 40.0, 45.0, 40.0
            content_w = page_w - (2 * margin_x)
            max_y = page_h - margin_bottom
            col_w = min(120.0, max(50.0, content_w / max(1, num_cols)))
            row_h = 22.0
            doc_out = fitz.open()
            page = doc_out.new_page(width=page_w, height=page_h)
            y = margin_top
            for row_idx, row in enumerate(rows):
                if y + row_h > max_y:
                    page = doc_out.new_page(width=page_w, height=page_h)
                    y = margin_top
                is_header = row_idx == 0
                for col_idx in range(num_cols):
                    val = str(row[col_idx]) if col_idx < len(row) and row[col_idx] is not None else ""
                    x0 = margin_x + col_idx * col_w
                    y0 = y
                    x1 = x0 + col_w
                    y1 = y0 + row_h
                    if is_header:
                        page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.7, 0.7, 0.8), fill=(0.2, 0.18, 0.35))
                        page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=9, color=(1, 1, 1))
                    else:
                        bg = (0.97, 0.97, 0.99) if row_idx % 2 == 0 else (1, 1, 1)
                        page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.85, 0.85, 0.88), fill=bg)
                        page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=8.5, color=(0.15, 0.15, 0.15))
                y += row_h
            buf = BytesIO()
            doc_out.save(buf, garbage=4, deflate=True, clean=True)
            doc_out.close()
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buf.getvalue(), "pdf")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def rtf(self, job, target):
        binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
        if has_binary(binary) or has_binary("soffice"):
            try:
                return self.libreoffice(job, target)
            except Exception:
                pass
        raw = read_upload(job.original).decode("utf-8", errors="replace")
        try:
            from striprtf.striprtf import rtf_to_text

            text = rtf_to_text(raw)
        except Exception:
            text = re.sub(r"\\[a-z0-9]+", " ", raw)
            text = re.sub(r"[{}\\]", "", text).strip()
        if target == "txt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", text.encode("utf-8"), "txt")
        if target == "docx":
            from docx import Document

            doc = Document()
            for line in text.splitlines():
                if line.strip():
                    doc.add_paragraph(line.strip())
            buf = BytesIO()
            doc.save(buf)
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.docx", buf.getvalue(), "docx")
        if target == "pdf":
            return text_to_pdf(text, f"{safe_stem(job.original_name)}-thuna.pdf")
        if target == "html":
            body = "<br>\n".join(html.escape(line) for line in text.splitlines() if line.strip())
            return ConversionResult(
                f"{safe_stem(job.original_name)}-thuna.html",
                f"<!doctype html><html><body>{body}</body></html>".encode("utf-8"),
                "html",
            )
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def ppt(self, job, target):
        binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
        if has_binary(binary) or has_binary("soffice"):
            try:
                return self.libreoffice(job, target)
            except Exception:
                pass
        raw = read_upload(job.original)
        text = self._extract_text_from_ole(raw)
        if not text:
            text = f"Presentation: {job.original_name}"
        if target == "txt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", text.encode("utf-8"), "txt")
        if target == "pdf":
            return text_to_pdf(text, f"{safe_stem(job.original_name)}-thuna.pdf")
        if target == "ppt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.ppt", raw, "ppt")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def libreoffice(self, job, target):
        binary = binary_setting("LIBREOFFICE_BINARY", "libreoffice")
        if not has_binary(binary):
            binary = "soffice"
        if not has_binary(binary):
            raise ConversionError("MISSING_ENGINE", "Office conversion engine is not installed on this server.")

        input_path = self.workdir / f"input.{job.input_format}"
        input_path.write_bytes(read_upload(job.original))
        command = [binary, "--headless", "--convert-to", target, "--outdir", str(self.workdir), str(input_path)]
        try:
            proc = subprocess.run(
                command,
                check=False,
                capture_output=True,
                timeout=setting("CONVERSION_TIMEOUT", 120),
            )
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("TIMEOUT", ERROR_MESSAGES["TIMEOUT"]) from exc
        if proc.returncode != 0:
            raise ConversionError("CONVERSION_FAILED", "Office conversion failed.")
        outputs = sorted(self.workdir.glob(f"*.{target}"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not outputs:
            raise ConversionError("CONVERSION_FAILED", "Office conversion did not create an output file.")
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", outputs[0].read_bytes(), target)

    def docx_to_pdf(self, job):
        try:
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "DOCX conversion requires python-docx.") from exc
        fitz = get_fitz()

        doc_in = Document(BytesIO(read_upload(job.original)))
        doc_out = fitz.open()
        page_w, page_h = 595.0, 842.0
        margin_x, margin_top, margin_bottom = 54.0, 54.0, 54.0
        content_w = page_w - (2 * margin_x)
        max_y = page_h - margin_bottom

        page = doc_out.new_page(width=page_w, height=page_h)
        y = margin_top

        def ensure_space(needed):
            nonlocal page, y
            if y + needed > max_y:
                page = doc_out.new_page(width=page_w, height=page_h)
                y = margin_top

        for element in doc_in.element.body:
            tag = element.tag.split("}")[-1]
            if tag == "p":
                p = Paragraph(element, doc_in)
                text = p.text.strip()
                if not text:
                    y += 8
                    continue

                style_name = (p.style.name or "").lower()
                if "title" in style_name:
                    fontsize, line_h, fontname = 22, 28, "helv"
                    ensure_space(line_h + 10)
                elif "heading 1" in style_name:
                    fontsize, line_h, fontname = 16, 22, "helv"
                    ensure_space(line_h + 8)
                elif "heading 2" in style_name:
                    fontsize, line_h, fontname = 13, 18, "helv"
                    ensure_space(line_h + 6)
                elif "heading 3" in style_name:
                    fontsize, line_h, fontname = 11.5, 16, "helv"
                    ensure_space(line_h + 4)
                else:
                    fontsize, line_h, fontname = 10, 14, "helv"

                is_bullet = "bullet" in style_name or "list" in style_name
                indent = 16 if is_bullet else 0
                avail_w = content_w - indent

                words = text.split(" ")
                lines = []
                curr_line = []
                for word in words:
                    test_line = " ".join(curr_line + [word])
                    if fitz.get_text_length(test_line, fontname=fontname, fontsize=fontsize) <= avail_w:
                        curr_line.append(word)
                    else:
                        if curr_line:
                            lines.append(" ".join(curr_line))
                            curr_line = [word]
                        else:
                            lines.append(word)
                            curr_line = []
                if curr_line:
                    lines.append(" ".join(curr_line))

                for i, line in enumerate(lines):
                    ensure_space(line_h)
                    if is_bullet and i == 0:
                        page.insert_text((margin_x + 4, y + fontsize), "•", fontname=fontname, fontsize=fontsize, color=(0.2, 0.2, 0.2))
                    page.insert_text((margin_x + indent, y + fontsize), line, fontname=fontname, fontsize=fontsize, color=(0.1, 0.1, 0.1))
                    y += line_h
                y += 4

            elif tag == "tbl":
                table = Table(element, doc_in)
                num_cols = len(table.columns) if table.columns else 1
                col_w = content_w / max(1, num_cols)
                row_h = 20
                for row_idx, row in enumerate(table.rows):
                    ensure_space(row_h + 4)
                    for col_idx, cell in enumerate(row.cells):
                        x0 = margin_x + col_idx * col_w
                        y0 = y
                        x1 = x0 + col_w
                        y1 = y0 + row_h
                        if row_idx == 0:
                            page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.8, 0.8, 0.85), fill=(0.92, 0.93, 0.96))
                        else:
                            page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.85, 0.85, 0.88))
                        cell_text = cell.text.strip()[:60]
                        page.insert_text((x0 + 4, y0 + 13), cell_text, fontname="helv", fontsize=8.5, color=(0.1, 0.1, 0.1))
                    y += row_h
                y += 8

        total_pages = doc_out.page_count
        for idx, p in enumerate(doc_out):
            p.insert_text((margin_x, page_h - 24), f"Page {idx + 1} of {total_pages}", fontname="helv", fontsize=8, color=(0.5, 0.5, 0.5))
            p.insert_text((page_w - margin_x - 60, page_h - 24), "THUNA PDF", fontname="helv", fontsize=8, color=(0.5, 0.5, 0.5))

        buffer = BytesIO()
        doc_out.save(buffer, garbage=4, deflate=True, clean=True)
        doc_out.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buffer.getvalue(), "pdf")

    def docx_to_txt(self, job):
        try:
            from docx import Document
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "DOCX to TXT requires python-docx.") from exc
        doc = Document(BytesIO(read_upload(job.original)))
        text = []
        for p in doc.paragraphs:
            if p.text.strip():
                text.append(p.text.strip())
        for t in doc.tables:
            for row in t.rows:
                row_text = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_text:
                    text.append(" | ".join(row_text))
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", "\n\n".join(text).encode("utf-8"), "txt")

    def docx_to_html(self, job):
        try:
            from docx import Document
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "DOCX to HTML requires python-docx.") from exc
        doc = Document(BytesIO(read_upload(job.original)))
        body_parts = []
        for p in doc.paragraphs:
            txt = html.escape(p.text.strip())
            if not txt:
                continue
            style = (p.style.name or "").lower()
            if "title" in style:
                body_parts.append(f"<h1>{txt}</h1>")
            elif "heading 1" in style:
                body_parts.append(f"<h2>{txt}</h2>")
            elif "heading 2" in style:
                body_parts.append(f"<h3>{txt}</h3>")
            else:
                body_parts.append(f"<p>{txt}</p>")
        for t in doc.tables:
            body_parts.append("<table border='1' cellpadding='6' cellspacing='0'>")
            for row in t.rows:
                body_parts.append("<tr>" + "".join(f"<td>{html.escape(c.text.strip())}</td>" for c in row.cells) + "</tr>")
            body_parts.append("</table>")
        html_content = (
            f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(job.original_name)}</title>"
            f"<style>body{{font-family:sans-serif;max-width:800px;margin:40px auto;line-height:1.6;padding:0 20px;}}"
            f"table{{border-collapse:collapse;width:100%;margin:20px 0;}}td,th{{border:1px solid #ddd;padding:8px;}}</style>"
            f"</head><body>{''.join(body_parts)}</body></html>"
        )
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.html", html_content.encode("utf-8"), "html")

    def xlsx_to_csv(self, job):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "XLSX conversion requires openpyxl.") from exc
        wb = load_workbook(BytesIO(read_upload(job.original)), data_only=True)
        ws = wb.active
        out = StringIO()
        writer = csv.writer(out, lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                writer.writerow([str(c) if c is not None else "" for c in row])
        wb.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.csv", out.getvalue().encode("utf-8"), "csv")

    def xlsx_to_pdf(self, job):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "XLSX conversion requires openpyxl.") from exc
        fitz = get_fitz()
        wb = load_workbook(BytesIO(read_upload(job.original)), data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        wb.close()
        if not rows:
            raise ConversionError("CONVERSION_FAILED", "Worksheet contains no readable rows.")

        num_cols = max(len(r) for r in rows)
        is_landscape = num_cols > 6
        page_w = 842.0 if is_landscape else 595.0
        page_h = 595.0 if is_landscape else 842.0
        margin_x, margin_top, margin_bottom = 40.0, 45.0, 40.0
        content_w = page_w - (2 * margin_x)
        max_y = page_h - margin_bottom

        col_w = min(120.0, max(50.0, content_w / max(1, num_cols)))
        row_h = 22.0

        doc_out = fitz.open()
        page = doc_out.new_page(width=page_w, height=page_h)
        y = margin_top

        sheet_title = getattr(ws, "title", "Sheet1")
        page.insert_text((margin_x, y), f"Sheet: {sheet_title}", fontname="helv", fontsize=13, color=(0.1, 0.1, 0.2))
        y += 20

        for row_idx, row in enumerate(rows):
            if y + row_h > max_y:
                page = doc_out.new_page(width=page_w, height=page_h)
                y = margin_top
            is_header = row_idx == 0
            for col_idx in range(num_cols):
                val = str(row[col_idx]) if col_idx < len(row) and row[col_idx] is not None else ""
                x0 = margin_x + col_idx * col_w
                y0 = y
                x1 = x0 + col_w
                y1 = y0 + row_h
                if is_header:
                    page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.7, 0.7, 0.8), fill=(0.2, 0.18, 0.35))
                    page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=9, color=(1, 1, 1))
                else:
                    bg = (0.97, 0.97, 0.99) if row_idx % 2 == 0 else (1, 1, 1)
                    page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.85, 0.85, 0.88), fill=bg)
                    page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=8.5, color=(0.15, 0.15, 0.15))
            y += row_h

        buffer = BytesIO()
        doc_out.save(buffer, garbage=4, deflate=True, clean=True)
        doc_out.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buffer.getvalue(), "pdf")

    def xlsx_to_txt(self, job):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "XLSX conversion requires openpyxl.") from exc
        wb = load_workbook(BytesIO(read_upload(job.original)), data_only=True)
        ws = wb.active
        lines = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
        wb.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", "\n".join(lines).encode("utf-8"), "txt")

    def xlsx_to_html(self, job):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "XLSX conversion requires openpyxl.") from exc
        wb = load_workbook(BytesIO(read_upload(job.original)), data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        wb.close()
        body = ["<table border='1' cellpadding='6' cellspacing='0'>"]
        for idx, row in enumerate(rows):
            tag = "th" if idx == 0 else "td"
            cells = "".join(f"<{tag}>{html.escape(str(c) if c is not None else '')}</{tag}>" for c in row)
            body.append(f"<tr>{cells}</tr>")
        body.append("</table>")
        html_content = (
            f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(job.original_name)}</title>"
            f"<style>body{{font-family:sans-serif;max-width:900px;margin:40px auto;padding:0 20px;}}"
            f"table{{border-collapse:collapse;width:100%;}}th{{background:#f0f2f5;text-align:left;}}td,th{{border:1px solid #ccc;padding:8px;}}</style>"
            f"</head><body><h1>{html.escape(getattr(ws, 'title', 'Sheet1'))}</h1>{''.join(body)}</body></html>"
        )
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.html", html_content.encode("utf-8"), "html")

    def pptx_to_pdf(self, job):
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "PPTX conversion requires python-pptx.") from exc
        fitz = get_fitz()
        prs = Presentation(BytesIO(read_upload(job.original)))
        doc_out = fitz.open()
        page_w, page_h = 842.0, 474.0
        for idx, slide in enumerate(prs.slides):
            page = doc_out.new_page(width=page_w, height=page_h)
            page.draw_rect(fitz.Rect(0, 0, page_w, page_h), fill=(0.98, 0.98, 1.0))
            page.draw_rect(fitz.Rect(0, 0, page_w, 6), fill=(0.4, 0.25, 0.8))

            y = 50.0
            if slide.shapes.title and slide.shapes.title.text.strip():
                page.insert_text((50, y), slide.shapes.title.text.strip()[:80], fontname="helv", fontsize=20, color=(0.1, 0.1, 0.25))
                y += 35

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = p.text.strip()
                        if txt:
                            indent = 20 if p.level > 0 else 0
                            page.insert_text((50 + indent, y), f"• {txt}" if p.level > 0 else txt, fontname="helv", fontsize=11, color=(0.2, 0.2, 0.2))
                            y += 18
                            if y > page_h - 40:
                                break
            page.insert_text((page_w - 90, page_h - 20), f"Slide {idx + 1}", fontname="helv", fontsize=9, color=(0.5, 0.5, 0.5))

        buffer = BytesIO()
        doc_out.save(buffer, garbage=4, deflate=True, clean=True)
        doc_out.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buffer.getvalue(), "pdf")

    def pptx_to_txt(self, job):
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "PPTX conversion requires python-pptx.") from exc
        prs = Presentation(BytesIO(read_upload(job.original)))
        parts = []
        for idx, slide in enumerate(prs.slides):
            slide_text = [f"--- Slide {idx + 1} ---"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        if p.text.strip():
                            slide_text.append(p.text.strip())
            parts.append("\n".join(slide_text))
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", "\n\n".join(parts).encode("utf-8"), "txt")

    def compress_docx(self, job, quality=80):
        original_data = read_upload(job.original)
        in_buf = BytesIO(original_data)
        out_buf = BytesIO()
        compressed_count = 0
        with ZipFile(in_buf, "r") as zin, ZipFile(out_buf, "w", compression=ZIP_DEFLATED, compresslevel=9) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                fname = item.filename.lower()
                if fname.startswith("word/media/") and any(fname.endswith(ext) for ext in [".png", ".jpg", ".jpeg"]):
                    try:
                        img = Image.open(BytesIO(data))
                        max_dim = 1600
                        if img.width > max_dim or img.height > max_dim:
                            img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)
                        img_buf = BytesIO()
                        fmt = "JPEG" if fname.endswith((".jpg", ".jpeg")) else "PNG"
                        if fmt == "JPEG":
                            if img.mode not in ("RGB", "L"):
                                img = img.convert("RGB")
                            img.save(img_buf, format=fmt, quality=int(quality), optimize=True)
                        else:
                            img.save(img_buf, format=fmt, optimize=True)
                        new_data = img_buf.getvalue()
                        if len(new_data) < len(data):
                            data = new_data
                            compressed_count += 1
                    except Exception:
                        pass
                zout.writestr(item, data)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.docx", out_buf.getvalue(), "docx")

    def txt(self, job, target):
        text = read_upload(job.original).decode("utf-8", errors="replace")
        if target in {"txt", "csv"}:
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", text.encode("utf-8"), target)
        if target == "docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise ConversionError("MISSING_ENGINE", "TXT to DOCX requires python-docx.") from exc
            document = Document()
            for line in text.splitlines() or [""]:
                document.add_paragraph(line)
            buffer = BytesIO()
            document.save(buffer)
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.docx", buffer.getvalue(), "docx")
        if target == "html":
            body = "<br>\n".join(html.escape(line) for line in text.splitlines())
            data = f"<!doctype html><html><body>{body}</body></html>".encode("utf-8")
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.html", data, "html")
        if target == "pdf":
            return text_to_pdf(text, f"{safe_stem(job.original_name)}-thuna.pdf")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def markdown(self, job, target):
        source = read_upload(job.original).decode("utf-8", errors="replace")
        plain = re.sub(r"(```.*?```|`[^`]*`)", "", source, flags=re.S)
        plain = re.sub(r"[#*_>\[\]()`-]+", "", plain)
        if target == "txt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", plain.encode("utf-8"), "txt")
        if target == "html":
            body = "<br>\n".join(html.escape(line) for line in source.splitlines())
            data = f"<!doctype html><html><body>{body}</body></html>".encode("utf-8")
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.html", data, "html")
        if target == "docx":
            return self.txt(job, "docx")
        if target == "pdf":
            return text_to_pdf(plain, f"{safe_stem(job.original_name)}-thuna.pdf")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def html(self, job, target):
        source = read_upload(job.original).decode("utf-8", errors="replace")
        text = re.sub(r"<(script|style).*?</\1>", "", source, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
        text = html.unescape(re.sub(r"\s+", " ", text)).strip()
        if target == "txt":
            return ConversionResult(f"{safe_stem(job.original_name)}-thuna.txt", text.encode("utf-8"), "txt")
        if target == "pdf":
            return text_to_pdf(text, f"{safe_stem(job.original_name)}-thuna.pdf")
        if target == "docx":
            return self.txt(job, "docx")
        raise ConversionError("UNSUPPORTED_FORMAT", ERROR_MESSAGES["UNSUPPORTED_FORMAT"])

    def csv_to_xlsx(self, job):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise ConversionError("MISSING_ENGINE", "CSV to XLSX requires openpyxl.") from exc
        text = read_upload(job.original).decode("utf-8", errors="replace")
        workbook = Workbook()
        sheet = workbook.active
        for row in csv.reader(StringIO(text)):
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.xlsx", buffer.getvalue(), "xlsx")

    def csv_to_pdf(self, job):
        fitz = get_fitz()
        text = read_upload(job.original).decode("utf-8", errors="replace")
        rows = list(csv.reader(StringIO(text)))
        if not rows:
            raise ConversionError("CONVERSION_FAILED", "CSV file contains no rows.")
        num_cols = max(len(r) for r in rows)
        is_landscape = num_cols > 6
        page_w = 842.0 if is_landscape else 595.0
        page_h = 595.0 if is_landscape else 842.0
        margin_x, margin_top, margin_bottom = 40.0, 45.0, 40.0
        content_w = page_w - (2 * margin_x)
        max_y = page_h - margin_bottom
        col_w = min(120.0, max(50.0, content_w / max(1, num_cols)))
        row_h = 22.0
        doc_out = fitz.open()
        page = doc_out.new_page(width=page_w, height=page_h)
        y = margin_top
        for row_idx, row in enumerate(rows):
            if y + row_h > max_y:
                page = doc_out.new_page(width=page_w, height=page_h)
                y = margin_top
            is_header = row_idx == 0
            for col_idx in range(num_cols):
                val = row[col_idx] if col_idx < len(row) else ""
                x0 = margin_x + col_idx * col_w
                y0 = y
                x1 = x0 + col_w
                y1 = y0 + row_h
                if is_header:
                    page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.7, 0.7, 0.8), fill=(0.2, 0.18, 0.35))
                    page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=9, color=(1, 1, 1))
                else:
                    bg = (0.97, 0.97, 0.99) if row_idx % 2 == 0 else (1, 1, 1)
                    page.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.85, 0.85, 0.88), fill=bg)
                    page.insert_text((x0 + 4, y0 + 15), val[:25], fontname="helv", fontsize=8.5, color=(0.15, 0.15, 0.15))
            y += row_h
        buffer = BytesIO()
        doc_out.save(buffer, garbage=4, deflate=True, clean=True)
        doc_out.close()
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.pdf", buffer.getvalue(), "pdf")


class MediaConverter:
    VIDEO_CODECS = {"mp4": "libx264", "mov": "libx264", "avi": "mpeg4", "mkv": "libx264", "webm": "libvpx-vp9"}
    AUDIO_CODECS = {"mp3": "libmp3lame", "wav": "pcm_s16le", "m4a": "aac", "aac": "aac", "flac": "flac", "ogg": "libvorbis"}

    def __init__(self, workdir):
        self.workdir = Path(workdir)

    def convert(self, job, source_format, target):
        ffmpeg = shutil.which(binary_setting("FFMPEG_BINARY", "ffmpeg")) or binary_setting("FFMPEG_BINARY", "ffmpeg")
        ffprobe = shutil.which(binary_setting("FFPROBE_BINARY", "ffprobe")) or binary_setting("FFPROBE_BINARY", "ffprobe")
        if not has_binary(ffmpeg) or not has_binary(ffprobe):
            raise ConversionError("MISSING_ENGINE", "FFmpeg and FFprobe are required for media conversion.")
        input_path = self.workdir / f"input.{source_format}"
        output_path = self.workdir / f"output.{target}"
        input_path.write_bytes(read_upload(job.original))
        info = self.probe(ffprobe, input_path)
        streams = info.get("streams", [])
        has_video = any(s.get("codec_type") == "video" for s in streams)
        has_audio = any(s.get("codec_type") == "audio" for s in streams)
        if target in VIDEO_EXTENSIONS and not has_video:
            raise ConversionError("NO_VIDEO_STREAM", ERROR_MESSAGES["NO_VIDEO_STREAM"])
        if target in AUDIO_EXTENSIONS and not has_audio:
            raise ConversionError("NO_AUDIO_STREAM", ERROR_MESSAGES["NO_AUDIO_STREAM"])

        command = [ffmpeg, "-y", "-i", str(input_path)]
        if target in VIDEO_EXTENSIONS:
            codec = self.VIDEO_CODECS.get(target, "libx264")
            command += ["-c:v", codec]
            if has_audio:
                if target == "webm":
                    command += ["-c:a", "libvorbis"]
                else:
                    command += ["-c:a", "aac"]
            else:
                command += ["-an"]
            if target == "mp4":
                command += ["-movflags", "+faststart", "-pix_fmt", "yuv420p"]
        else:
            codec = self.AUDIO_CODECS.get(target, "libmp3lame")
            command += ["-vn", "-c:a", codec]
            if target in {"mp3", "m4a", "aac"}:
                command += ["-b:a", "192k"]
        command.append(str(output_path))
        self.run(command)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{target}", output_path.read_bytes(), target)

    def compress(self, job, source_format, quality=80):
        ffmpeg = shutil.which(binary_setting("FFMPEG_BINARY", "ffmpeg")) or binary_setting("FFMPEG_BINARY", "ffmpeg")
        ffprobe = shutil.which(binary_setting("FFPROBE_BINARY", "ffprobe")) or binary_setting("FFPROBE_BINARY", "ffprobe")
        if not has_binary(ffmpeg) or not has_binary(ffprobe):
            raise ConversionError("MISSING_ENGINE", "FFmpeg and FFprobe are required for media conversion.")
        input_path = self.workdir / f"input.{source_format}"
        output_path = self.workdir / f"output.{source_format}"
        input_path.write_bytes(read_upload(job.original))
        info = self.probe(ffprobe, input_path)
        streams = info.get("streams", [])
        has_video = any(s.get("codec_type") == "video" for s in streams)
        has_audio = any(s.get("codec_type") == "audio" for s in streams)

        q = max(10, min(100, int(quality or 80)))
        command = [ffmpeg, "-y", "-i", str(input_path)]
        if source_format in VIDEO_EXTENSIONS:
            crf = int(18 + ((100 - q) / 100.0) * 18)
            codec = self.VIDEO_CODECS.get(source_format, "libx264")
            command += ["-c:v", codec, "-crf", str(crf), "-preset", "fast"]
            if has_audio:
                audio_bitrate = "128k" if q >= 75 else ("96k" if q >= 45 else "64k")
                if source_format == "webm":
                    command += ["-c:a", "libvorbis", "-b:a", audio_bitrate]
                else:
                    command += ["-c:a", "aac", "-b:a", audio_bitrate]
            else:
                command += ["-an"]
            if source_format == "mp4":
                command += ["-movflags", "+faststart", "-pix_fmt", "yuv420p"]
        elif source_format in AUDIO_EXTENSIONS:
            audio_bitrate = "192k" if q >= 85 else ("128k" if q >= 65 else ("96k" if q >= 40 else "64k"))
            codec = self.AUDIO_CODECS.get(source_format, "libmp3lame")
            command += ["-vn", "-c:a", codec]
            if codec not in {"flac", "pcm_s16le"}:
                command += ["-b:a", audio_bitrate]
        else:
            raise ConversionError("UNSUPPORTED_FORMAT", f"Compression is not supported for {source_format.upper()}.")

        command.append(str(output_path))
        self.run(command)
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.{source_format}", output_path.read_bytes(), source_format)

    def probe(self, ffprobe, path):
        import json

        command = [ffprobe, "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)]
        try:
            proc = subprocess.run(command, check=False, capture_output=True, timeout=setting("CONVERSION_TIMEOUT", 120))
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("TIMEOUT", ERROR_MESSAGES["TIMEOUT"]) from exc
        if proc.returncode != 0:
            ffmpeg = shutil.which(binary_setting("FFMPEG_BINARY", "ffmpeg")) or binary_setting("FFMPEG_BINARY", "ffmpeg")
            if has_binary(ffmpeg):
                try:
                    f_proc = subprocess.run([ffmpeg, "-i", str(path)], check=False, capture_output=True, text=True, errors="replace", timeout=30)
                    combined_out = (f_proc.stdout or "") + (f_proc.stderr or "")
                    streams = []
                    if "Video:" in combined_out:
                        streams.append({"codec_type": "video"})
                    if "Audio:" in combined_out:
                        streams.append({"codec_type": "audio"})
                    if streams:
                        return {"streams": streams, "format": {}}
                except Exception:
                    pass
            raise ConversionError("CORRUPTED_FILE", "FFprobe could not read this media file.")
        try:
            return json.loads(proc.stdout.decode("utf-8", errors="replace"))
        except Exception:
            return {"streams": [{"codec_type": "video"}, {"codec_type": "audio"}]}

    def run(self, command):
        try:
            proc = subprocess.run(command, check=False, capture_output=True, timeout=setting("CONVERSION_TIMEOUT", 120))
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("TIMEOUT", ERROR_MESSAGES["TIMEOUT"]) from exc
        if proc.returncode != 0:
            raise ConversionError("CONVERSION_FAILED", "FFmpeg conversion failed.")


class ArchiveConverter:
    def __init__(self, workdir):
        self.workdir = Path(workdir)

    def create_zip(self, uploads, archive_name="download"):
        if not uploads:
            raise ConversionError("INVALID_FILE", "Choose at least one file to create a ZIP.")
        buffer = BytesIO()
        used = set()
        with ZipFile(buffer, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
            for uploaded in uploads:
                arcname = unique_arcname(safe_archive_name(uploaded.name), used)
                archive.writestr(arcname, read_upload(uploaded))
        return ConversionResult(f"{safe_stem(archive_name)}.zip", buffer.getvalue(), "zip", {"files": len(uploads)})

    def recompress_zip(self, job):
        original_data = read_upload(job.original)
        self.inspect_zip(original_data)
        buffer = BytesIO()
        with ZipFile(BytesIO(original_data), "r") as source, ZipFile(buffer, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
            for entry in source.infolist():
                if entry.is_dir():
                    target.writestr(entry.filename.rstrip("/") + "/", b"")
                    continue
                target.writestr(safe_archive_name(entry.filename), source.read(entry.filename))
        recompressed = buffer.getvalue()
        if len(recompressed) >= len(original_data):
            raise ConversionError(
                "CONVERSION_FAILED",
                "This ZIP is already efficiently compressed. Re-compressing it would not reduce the size.",
            )
        return ConversionResult(f"{safe_stem(job.original_name)}-thuna.zip", recompressed, "zip")

    def extract_zip(self, job):
        original_data = read_upload(job.original)
        self.inspect_zip(original_data)
        extract_dir = self.workdir / "extracted"
        extract_dir.mkdir()
        with ZipFile(BytesIO(original_data), "r") as archive:
            for entry in archive.infolist():
                if entry.is_dir():
                    continue
                destination = safe_extract_path(extract_dir, entry.filename)
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(archive.read(entry.filename))
        files = [p for p in extract_dir.rglob("*") if p.is_file()]
        return self.create_zip([LocalUpload(path) for path in files], f"{safe_stem(job.original_name)}-extracted")

    def inspect_zip(self, data):
        max_archive = setting("MAX_ARCHIVE_SIZE", 200 * 1024 * 1024)
        max_extracted = setting("MAX_EXTRACTED_SIZE", 500 * 1024 * 1024)
        max_files = setting("MAX_ARCHIVE_FILES", 2000)
        if len(data) > max_archive:
            raise ConversionError("ARCHIVE_TOO_LARGE", ERROR_MESSAGES["ARCHIVE_TOO_LARGE"])
        try:
            with ZipFile(BytesIO(data), "r") as archive:
                infos = archive.infolist()
                if len(infos) > max_files:
                    raise ConversionError("ARCHIVE_TOO_LARGE", "The archive contains too many files.")
                total = sum(item.file_size for item in infos)
                if total > max_extracted:
                    raise ConversionError("ARCHIVE_TOO_LARGE", ERROR_MESSAGES["ARCHIVE_TOO_LARGE"])
                for item in infos:
                    safe_extract_path(self.workdir, item.filename)
                if archive.testzip() is not None:
                    raise ConversionError("CORRUPTED_FILE", "This ZIP file appears corrupted.")
        except BadZipFile as exc:
            raise ConversionError("CORRUPTED_FILE", "This ZIP file appears corrupted.") from exc


class LocalUpload:
    def __init__(self, path):
        self.path = Path(path)
        self.name = self.path.name
        self.size = self.path.stat().st_size

    def open(self, mode="rb"):
        return self.path.open(mode)

    def seek(self, *_args):
        return None

    def read(self):
        return self.path.read_bytes()


def safe_archive_name(name):
    normalized = PurePosixPath(str(name).replace("\\", "/"))
    parts = [part for part in normalized.parts if part not in {"", ".", ".."} and not re.match(r"^[A-Za-z]:$", part)]
    if not parts:
        return f"file-{uuid.uuid4().hex}"
    return "/".join(re.sub(r"[^A-Za-z0-9._ -]+", "_", part).strip() or "file" for part in parts)


def unique_arcname(name, used):
    base = PurePosixPath(name)
    stem = base.stem or "file"
    suffix = base.suffix
    parent = "" if str(base.parent) == "." else str(base.parent) + "/"
    candidate = f"{parent}{stem}{suffix}"
    index = 2
    while candidate in used:
        candidate = f"{parent}{stem}-{index}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def safe_extract_path(base_dir, member_name):
    member_name = member_name.replace("\\", "/")
    raw_parts = PurePosixPath(member_name).parts
    if (
        member_name.startswith("/")
        or re.match(r"^[A-Za-z]:", member_name)
        or "\x00" in member_name
        or ".." in raw_parts
    ):
        raise ConversionError("UNSAFE_ARCHIVE", ERROR_MESSAGES["UNSAFE_ARCHIVE"])
    destination = (Path(base_dir) / safe_archive_name(member_name)).resolve()
    if not str(destination).startswith(str(Path(base_dir).resolve())):
        raise ConversionError("UNSAFE_ARCHIVE", ERROR_MESSAGES["UNSAFE_ARCHIVE"])
    return destination


def text_to_pdf(text, name):
    fitz = get_fitz()
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    y = 50
    for line in (text or "").splitlines() or [""]:
        for chunk in [line[i : i + 90] for i in range(0, len(line), 90)] or [""]:
            if y > 790:
                page = doc.new_page(width=595, height=842)
                y = 50
            page.insert_text((50, y), chunk, fontsize=11)
            y += 15
    buffer = BytesIO()
    doc.save(buffer, garbage=4, deflate=True, clean=True)
    doc.close()
    return ConversionResult(name, buffer.getvalue(), "pdf")


def validate_output(data, output_format):
    if not data:
        raise ConversionError("OUTPUT_VALIDATION_FAILED", ERROR_MESSAGES["OUTPUT_VALIDATION_FAILED"])
    fmt = normalize_ext(output_format)
    try:
        if fmt in IMAGE_OUTPUT_FORMATS and fmt != "pdf":
            with Image.open(BytesIO(data)) as image:
                image.verify()
        elif fmt == "pdf":
            fitz = get_fitz()
            doc = fitz.open(stream=data, filetype="pdf")
            if doc.page_count < 1:
                raise ConversionError("OUTPUT_VALIDATION_FAILED", "The output PDF has no pages.")
            doc.close()
        elif fmt == "zip":
            with ZipFile(BytesIO(data), "r") as archive:
                if archive.testzip() is not None:
                    raise ConversionError("OUTPUT_VALIDATION_FAILED", "The output ZIP is invalid.")
        elif fmt == "docx":
            from docx import Document

            Document(BytesIO(data))
        elif fmt == "xlsx":
            from openpyxl import load_workbook

            load_workbook(BytesIO(data), read_only=True).close()
        elif fmt in {"txt", "csv", "html"}:
            data.decode("utf-8")
        elif fmt in VIDEO_EXTENSIONS | AUDIO_EXTENSIONS:
            if len(data) < 128:
                raise ConversionError("OUTPUT_VALIDATION_FAILED", ERROR_MESSAGES["OUTPUT_VALIDATION_FAILED"])
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("OUTPUT_VALIDATION_FAILED", ERROR_MESSAGES["OUTPUT_VALIDATION_FAILED"]) from exc


def save_success(job, result):
    job.result_name = result.name
    job.result.save(result.name, ContentFile(result.data), save=False)
    job.result_size = job.result.size
    job.output_path = job.result.name
    job.compression_ratio = (job.result_size / job.original_size) if job.original_size else None
    job.status = ProcessingJob.Status.SUCCESS
    job.error_code = ""
    job.error_message = ""
    job.completed_at = timezone.now()
    job.save(
        update_fields=[
            "result",
            "result_name",
            "result_size",
            "output_path",
            "compression_ratio",
            "status",
            "error_code",
            "error_message",
            "completed_at",
        ]
    )


def save_failure(job, exc):
    if isinstance(exc, ConversionError):
        code = exc.code
        message = str(exc) or ERROR_MESSAGES.get(code, ERROR_MESSAGES["CONVERSION_FAILED"])
    else:
        code = "CONVERSION_FAILED"
        message = ERROR_MESSAGES[code]
    job.status = ProcessingJob.Status.FAILED
    job.error_code = code
    job.error_message = message[:255]
    job.completed_at = timezone.now()
    job.save(update_fields=["status", "error_code", "error_message", "completed_at"])


def process_convert(job, output_format, quality):
    return ConversionRouter().convert(job, output_format, quality)


def process_image(job, output_format, quality):
    category, source = detect_file(job.original)
    target = normalize_ext(output_format) or source
    if category != "image":
        raise ConversionError("INVALID_FILE", "This operation requires an image file.")
    job.input_format = source
    job.output_format = target
    job.save(update_fields=["input_format", "output_format"])
    return ConversionRouter().convert(job, target, quality)


def process_pdf(job, quality=80):
    return ConversionRouter().convert(job, "pdf", quality)


def process_zip(job):
    return ConversionRouter().convert(job, "zip", 80)


def process_merge(job, files, output_format, quality):
    if output_format and normalize_ext(output_format) != "pdf":
        raise ConversionError("UNSUPPORTED_FORMAT", "Merged files are exported as PDF.")
    fitz = get_fitz()
    merged = fitz.open()
    try:
        for uploaded in files:
            category, _fmt = detect_file(uploaded)
            if category == "pdf":
                source = fitz.open(stream=read_upload(uploaded), filetype="pdf")
                merged.insert_pdf(source)
                source.close()
            elif category == "image":
                image = Image.open(uploaded)
                image = ImageOps.exif_transpose(image)
                if image.mode not in ("RGB", "L"):
                    image = image.convert("RGB")
                buffer = BytesIO()
                image.save(buffer, format="PDF", quality=int(quality))
                source = fitz.open(stream=buffer.getvalue(), filetype="pdf")
                merged.insert_pdf(source)
                source.close()
            elif category == "office" and _fmt == "docx":
                temp_job = ProcessingJob(original=uploaded, original_name=uploaded.name, input_format="docx")
                pdf_res = DocumentConverter(tempfile.gettempdir()).docx_to_pdf(temp_job)
                source = fitz.open(stream=pdf_res.data, filetype="pdf")
                merged.insert_pdf(source)
                source.close()
            elif category == "text":
                text = read_upload(uploaded).decode("utf-8", errors="replace")
                pdf_res = text_to_pdf(text, "text.pdf")
                source = fitz.open(stream=pdf_res.data, filetype="pdf")
                merged.insert_pdf(source)
                source.close()
            else:
                raise ConversionError(
                    "UNSUPPORTED_FORMAT",
                    f"File '{uploaded.name}' ({_fmt.upper()}) cannot be merged. Supported: PDF, Images, DOCX, TXT.",
                )
        if merged.page_count == 0:
            raise ConversionError("CONVERSION_FAILED", "No mergeable pages were found.")
        buffer = BytesIO()
        merged.save(buffer, garbage=4, deflate=True, clean=True)
        result = ConversionResult("thuna-merged.pdf", buffer.getvalue(), "pdf")
        validate_output(result.data, result.output_format)
        save_success(job, result)
        return result
    finally:
        merged.close()
